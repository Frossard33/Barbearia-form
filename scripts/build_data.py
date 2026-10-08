#!/usr/bin/env python3
"""Converte a planilha de inadimplência em js/data.js.

Uso: python3 scripts/build_data.py caminho/da/planilha.xlsx [AAAA-MM-DD da data do relatório]
"""
import sys, json, datetime as dt, collections, re
import openpyxl

CAT = {
    'JURIDICO': 'Judicial', 'JUDICIAL': 'Judicial',
    'FROTA': 'Frota',
    'TESOURARIA': 'Tesouraria', 'FISCAL ITINERANTE': 'Tesouraria',
    'COBRANCA EXTRA': 'Cobrança Extra',
    'PERMUTA': 'Permuta', 'GRUPO PERMUTA': 'Permuta',
    'ORGAO PUBLICO': 'Órgão Público', 'ORGÃO PUBLICO': 'Órgão Público',
    'AUDITORIA': 'Auditoria',
    'EM ABERTO': 'Cobrança Administrativa', 'BOLETO': 'Cobrança Administrativa',
    'KONCILLI': 'Mercado Livre', 'MERCADO LIVRE': 'Mercado Livre',
    'PNEUSTORE': 'Pneustore', 'PNEUTORE': 'Pneustore',
    'DIVIDA ANTIGA': 'Dívida Antiga',
    'PENDENCIAS DIVERSA': 'Pendências Diversas', 'PENDENCIAS DIVERSAS': 'Pendências Diversas', 'DIVERSAS': 'Pendências Diversas',
    'VANESSA': 'Conciliação', 'MARLA': 'Conciliação',
}

def iso(v):
    if isinstance(v, dt.datetime): return v.date().isoformat()
    if isinstance(v, str) and re.match(r'\d\d/\d\d/\d{4}', v):
        d, m, y = v[:10].split('/'); return f'{y}-{m}-{d}'
    return None

def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else 0.0

def main(path, ref):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb['Base de dados'].iter_rows(values_only=True))
    hdr = list(rows[0])
    ix = {h: i for i, h in enumerate(hdr)}

    # Pagamentos já realizados (colunas "Data de Pagamento"/"Valor" das abas de cobrança)
    pagos = {}
    for n in wb.sheetnames:
        if not n.startswith('Cobrança Adm'): continue
        w = list(wb[n].iter_rows(values_only=True))
        if not w or w[0][0] != 'Filial Código': continue
        h = list(w[0])
        if 'Data de Pagamento' not in h: continue
        ip = h.index('Data de Pagamento')
        for r in w[1:]:
            if r[ip] is not None:
                pagos[(str(r[0]), str(r[2]))] = (iso(r[ip]), num(r[ip + 1]))

    out, seen = [], collections.Counter()
    for r in rows[1:]:
        if not r[0] or str(r[0]) in ('Filial Código',): continue
        tipo = 'R' if r[ix['Pagar / Receber Descrição']] == 'RECEBER' else 'P'
        st = (r[ix['STATUS']] or '').strip().upper()
        cat = CAT.get(st, 'Cobrança Administrativa' if tipo == 'R' else 'Em aberto')
        if tipo == 'P' and st not in ('VANESSA', 'MARLA'): cat = 'Em aberto'
        f = str(r[ix['Filial Código']]).lstrip('0') or '0'
        tid = str(r[ix['Titulo Código']])
        base = f'{f}-{tid}-{r[ix["Titulo Parcela"]]}'
        seen[base] += 1
        t = {
            'k': base if seen[base] == 1 else f'{base}#{seen[base]}',
            'f': int(f), 't': tipo, 'id': tid,
            'c': str(r[ix['Cliente / Fornecedor Código']] or ''),
            'n': (r[ix['Cliente / Fornecedor Nome']] or '').strip(),
            'v': iso(r[ix['Data de Vencimento']]),
            'm': iso(r[ix['Data de Movimento']]),
            'val': num(r[ix['Título Valor']]), 'sal': num(r[ix['Título Saldo']]),
            'cat': cat,
            'cond': (r[ix['Condição Pagamento']] or '').strip(),
            'par': str(r[ix['Titulo Parcela']] or ''),
            'obs': (r[ix['Observacao']] or '')[:160] if isinstance(r[ix['Observacao']], str) else '',
        }
        pg = pagos.get((f, tid))
        if pg: t['pg'], t['pgv'] = pg
        out.append(t)

    data = {'ref': ref, 'gerado': dt.date.today().isoformat(), 'titulos': out}
    with open('js/data.js', 'w', encoding='utf-8') as fh:
        fh.write('window.DADOS = ' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(len(out), 'títulos;', len(pagos), 'pagos')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else '2026-10-06')
